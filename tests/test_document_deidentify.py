import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from longgate.document_deidentify import deidentify_file_copy
from longgate.utils import sha256_file


def test_html_preserves_dom_and_shared_identifier_mapping(tmp_path: Path):
    source = tmp_path / "page.html"
    source.write_text(
        """<!doctype html>
<html><head><title>Case</title></head>
<body>
<p>Contact person@example.com</p>
<a href="mailto:person@example.com" aria-label="person@example.com">Email</a>
<table><tr><td>192.0.2.1</td></tr></table>
<script>const hidden = "script@example.com";</script>
</body></html>""",
        encoding="utf-8",
    )
    before = sha256_file(source)

    result = deidentify_file_copy(source)

    output = Path(result.output_path)
    text = output.read_text(encoding="utf-8")
    assert "<table>" in text
    assert text.count("[EMAIL_001]") == 3
    assert "[IP_ADDRESS_001]" in text
    assert "script@example.com" in text
    assert result.status == "LOCAL_ONLY"
    assert result.release_allowed is False
    assert sha256_file(source) == before

    audit = json.loads(Path(result.audit_path).read_text(encoding="utf-8"))
    assert audit["processed_units"]["visible_text_nodes"] >= 1
    assert audit["unprocessed_regions"]["comments_doctype_or_ignored_text_nodes"] >= 1
    assert audit["remaining_direct_pii_hits"] >= 1
    assert "person@example.com" not in json.dumps(audit)


def test_xlsx_processes_all_sheets_comments_links_and_properties(tmp_path: Path):
    source = tmp_path / "study.xlsx"
    workbook = Workbook()
    visible = workbook.active
    visible.title = "Visible"
    visible["A1"] = "person@example.com"
    visible["A2"] = "person@example.com"
    visible["A3"] = '="formula@example.com"'
    visible["B1"] = "Click"
    visible["B1"].hyperlink = "mailto:person@example.com"
    visible["C1"].comment = Comment("person@example.com", "author@example.com")

    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "person@example.com"
    workbook.properties.creator = "owner@example.com"
    workbook.save(source)
    before = sha256_file(source)

    result = deidentify_file_copy(source)

    output = Path(result.output_path)
    rewritten = load_workbook(output, data_only=False, keep_links=True)
    assert rewritten["Visible"]["A1"].value == "[EMAIL_001]"
    assert rewritten["Visible"]["A2"].value == "[EMAIL_001]"
    assert rewritten["Hidden"]["A1"].value == "[EMAIL_001]"
    assert rewritten["Visible"]["A3"].value == '="formula@example.com"'
    assert rewritten["Visible"]["B1"].hyperlink.target == "mailto:[EMAIL_001]"
    assert rewritten["Visible"]["C1"].comment.text == "[EMAIL_001]"
    assert rewritten["Visible"]["C1"].comment.author == "[EMAIL_002]"
    assert rewritten.properties.creator == "[EMAIL_003]"
    assert result.status == "LOCAL_ONLY"
    assert sha256_file(source) == before

    audit = json.loads(Path(result.audit_path).read_text(encoding="utf-8"))
    assert audit["processed_units"]["worksheets"] == 2
    assert audit["unprocessed_regions"]["formula_cells"] >= 1
    assert audit["remaining_direct_pii_hits"] >= 1
    payload = json.dumps(audit)
    assert "formula@example.com" not in payload
    assert "person@example.com" not in payload


def test_xlsx_without_unprocessed_direct_pii_reaches_manual_review(tmp_path: Path):
    source = tmp_path / "clean-formula.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "person@example.com"
    workbook.active["A2"] = "=1+1"
    workbook.save(source)

    result = deidentify_file_copy(source)

    assert result.status == "MANUAL_REVIEW_REQUIRED"
    assert result.release_allowed is False
