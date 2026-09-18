import json
import os
from pathlib import Path

import pytest
from docx import Document

from longgate.batch_deidentify import deidentify_batch


def test_batch_uses_consistent_cross_file_placeholders_without_leaking_state(tmp_path: Path):
    source = tmp_path / "private-input"
    output = tmp_path / "private-output"
    source.mkdir()
    (source / "alpha.txt").write_text(
        "Contact person@example.com and second@example.com",
        encoding="utf-8",
    )
    (source / "beta.md").write_text(
        "# Follow-up\nperson@example.com",
        encoding="utf-8",
    )

    result = deidentify_batch(source, output)

    alpha = (output / "alpha.txt").read_text(encoding="utf-8")
    beta = (output / "beta.md").read_text(encoding="utf-8")
    assert "[EMAIL_001]" in alpha
    assert "[EMAIL_002]" in alpha
    assert "[EMAIL_001]" in beta
    assert result.processed_files == 2
    assert result.resumed_files == 0
    assert result.release_allowed is False

    state_text = (output / ".longgate-batch-state.json").read_text(encoding="utf-8")
    assert "person@example.com" not in state_text
    assert "second@example.com" not in state_text
    assert "alpha.txt" not in state_text
    assert "beta.md" not in state_text
    state = json.loads(state_text)
    assert len(state["completed"]) == 2

    key_path = output / ".longgate-batch.key"
    assert len(key_path.read_bytes()) == 32
    if os.name != "nt":
        assert key_path.stat().st_mode & 0o077 == 0


def test_batch_resume_skips_verified_unchanged_files(monkeypatch, tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()
    (source / "a.txt").write_text("person@example.com", encoding="utf-8")
    (source / "b.txt").write_text("other@example.com", encoding="utf-8")
    first = deidentify_batch(source, output)
    assert first.processed_files == 2

    monkeypatch.setattr(
        "longgate.batch_deidentify.deidentify_file_copy",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("unchanged verified files must not be reprocessed")
        ),
    )
    resumed = deidentify_batch(source, output, resume=True)
    assert resumed.processed_files == 0
    assert resumed.resumed_files == 2


def test_batch_resume_reprocesses_only_changed_input(tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()
    (source / "a.txt").write_text("person@example.com", encoding="utf-8")
    (source / "b.txt").write_text("other@example.com", encoding="utf-8")
    deidentify_batch(source, output)

    (source / "b.txt").write_text(
        "other@example.com and third@example.com",
        encoding="utf-8",
    )
    resumed = deidentify_batch(source, output, resume=True)
    assert resumed.processed_files == 1
    assert resumed.resumed_files == 1
    assert "[EMAIL_003]" in (output / "b.txt").read_text(encoding="utf-8")


def test_batch_checkpoint_survives_failure_and_resumes(monkeypatch, tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()
    (source / "a.txt").write_text("person@example.com", encoding="utf-8")
    (source / "b.docx").write_bytes(b"broken-docx")

    with pytest.raises(ValueError, match="valid DOCX"):
        deidentify_batch(source, output)

    state = json.loads(
        (output / ".longgate-batch-state.json").read_text(encoding="utf-8")
    )
    assert len(state["completed"]) == 1

    document = Document()
    document.add_paragraph("person@example.com")
    document.save(source / "b.docx")

    resumed = deidentify_batch(source, output, resume=True)
    assert resumed.resumed_files == 1
    assert resumed.processed_files == 1

    rewritten = Document(output / "b.docx")
    assert rewritten.paragraphs[0].text == "[EMAIL_001]"


def test_batch_rejects_output_inside_input(tmp_path: Path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "a.txt").write_text("person@example.com", encoding="utf-8")

    with pytest.raises(ValueError, match="outside the input"):
        deidentify_batch(source, source / "out")


def test_batch_rejects_nonempty_fresh_output(tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()
    output.mkdir()
    (source / "a.txt").write_text("person@example.com", encoding="utf-8")
    (output / "existing.txt").write_text("do not overwrite", encoding="utf-8")

    with pytest.raises(ValueError, match="must be empty"):
        deidentify_batch(source, output)


def test_batch_resume_requires_complete_checkpoint_pair(tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()
    output.mkdir()
    (source / "a.txt").write_text("person@example.com", encoding="utf-8")

    with pytest.raises(ValueError, match="checkpoint/key pair"):
        deidentify_batch(source, output, resume=True)


def test_batch_resume_rejects_tampered_checkpoint(tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()
    (source / "a.txt").write_text("person@example.com", encoding="utf-8")
    deidentify_batch(source, output)

    state_path = output / ".longgate-batch-state.json"
    state = json.loads(state_path.read_text(encoding="utf-8"))
    record = next(iter(state["completed"].values()))
    record["output_sha256"] = "0" * 64
    state_path.write_text(json.dumps(state), encoding="utf-8")

    with pytest.raises(ValueError, match="authentication failed"):
        deidentify_batch(source, output, resume=True)


def test_batch_resume_rejects_wrong_local_key(tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()
    (source / "a.txt").write_text("person@example.com", encoding="utf-8")
    deidentify_batch(source, output)

    (output / ".longgate-batch.key").write_bytes(b"x" * 32)
    with pytest.raises(ValueError, match="authentication failed"):
        deidentify_batch(source, output, resume=True)


def test_batch_empty_input_creates_no_state(tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()

    with pytest.raises(ValueError, match="No supported files"):
        deidentify_batch(source, output)

    assert not output.exists()


def test_batch_rejects_case_colliding_paths(tmp_path: Path):
    source = tmp_path / "source"
    output = tmp_path / "output"
    source.mkdir()
    (source / "A.txt").write_text("one@example.com", encoding="utf-8")
    (source / "a.txt").write_text("two@example.com", encoding="utf-8")

    with pytest.raises(ValueError, match="case-colliding"):
        deidentify_batch(source, output)

    assert not output.exists()
