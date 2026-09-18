import json
from pathlib import Path

from docx import Document

from longgate.document_deidentify import deidentify_file_copy
from longgate.entity_assist import EntityAssistResult, validate_entity_payload
from longgate.format_deidentify import AssistedEntityLiteral, DirectIdentifierMapper


class FakeDetector:
    def __init__(
        self,
        candidates: list[AssistedEntityLiteral],
        *,
        rejected: int = 0,
        model_path: Path | None = None,
    ) -> None:
        self._candidates = candidates
        self._rejected = rejected
        self.model_path = model_path or Path("fake.gguf")

    def detect(self, source: str, *, max_tokens: int = 768) -> EntityAssistResult:
        accepted = [
            candidate for candidate in self._candidates if candidate.literal in source
        ]
        by_entity: dict[str, int] = {}
        for candidate in accepted:
            by_entity[candidate.entity] = by_entity.get(candidate.entity, 0) + 1
        return EntityAssistResult(
            candidates=accepted,
            accepted_by_entity=by_entity,
            rejected_candidates=self._rejected,
            model_file=self.model_path.name,
        )


def test_validate_entity_payload_accepts_exact_multilingual_literals():
    source = "张三在北京交通大学工作，2026年9月18日在北京市参加项目星河。"
    payload = json.dumps(
        {
            "entities": [
                {"type": "PERSON", "literal": "张三"},
                {"type": "ORGANIZATION", "literal": "北京交通大学"},
                {"type": "DATE", "literal": "2026年9月18日"},
                {"type": "LOCATION", "literal": "北京市"},
                {"type": "PROJECT", "literal": "项目星河"},
            ]
        },
        ensure_ascii=False,
    )

    accepted, rejected = validate_entity_payload(source, payload)

    assert rejected == 0
    assert {(item.entity, item.literal) for item in accepted} == {
        ("PERSON", "张三"),
        ("ORGANIZATION", "北京交通大学"),
        ("DATE", "2026年9月18日"),
        ("LOCATION", "北京市"),
        ("PROJECT", "项目星河"),
    }


def test_validate_entity_payload_rejects_hallucinated_or_unsupported_candidates():
    source = "张三在北京工作。"
    payload = json.dumps(
        {
            "entities": [
                {"type": "PERSON", "literal": "李四"},
                {"type": "SECRET", "literal": "张三"},
                {"type": "PERSON", "literal": "[PERSON_001]"},
            ]
        },
        ensure_ascii=False,
    )

    accepted, rejected = validate_entity_payload(source, payload)

    assert accepted == []
    assert rejected == 3


def test_deterministic_direct_pii_wins_over_overlapping_semantic_candidate():
    mapper = DirectIdentifierMapper()
    mapper.register_assisted_literals(
        [
            AssistedEntityLiteral(
                entity="QUASI_IDENTIFIER",
                literal="Contact person@example.com",
            )
        ]
    )

    transformed = mapper.replace("Contact person@example.com")

    assert transformed == "Contact [EMAIL_001]"
    assert mapper.entity_counts == {"EMAIL": 1}


def test_text_entity_assist_replaces_exact_literals_without_free_form_rewrite(
    tmp_path: Path,
):
    source = tmp_path / "interview.md"
    source.write_text(
        "# 访谈\n张三在北京交通大学工作，邮箱 zhang@example.com。\n",
        encoding="utf-8",
    )
    detector = FakeDetector(
        [
            AssistedEntityLiteral("PERSON", "张三"),
            AssistedEntityLiteral("ORGANIZATION", "北京交通大学"),
        ]
    )

    result = deidentify_file_copy(source, entity_detector=detector)

    output = Path(result.output_path).read_text(encoding="utf-8")
    assert output.startswith("# 访谈\n")
    assert "[PERSON_001]" in output
    assert "[ORGANIZATION_001]" in output
    assert "[EMAIL_001]" in output
    assert "张三" not in output
    assert "北京交通大学" not in output

    audit = json.loads(Path(result.audit_path).read_text(encoding="utf-8"))
    serialized = json.dumps(audit, ensure_ascii=False)
    assert "张三" not in serialized
    assert "北京交通大学" not in serialized
    assert audit["entity_assist"]["accepted_candidates"] == 2
    assert audit["entity_assist"]["free_form_rewrite"] is False
    assert result.status == "MANUAL_REVIEW_REQUIRED"


def test_rejected_semantic_candidate_forces_local_only(tmp_path: Path):
    source = tmp_path / "interview.txt"
    source.write_text("张三参加了研究。", encoding="utf-8")
    detector = FakeDetector(
        [AssistedEntityLiteral("PERSON", "张三")],
        rejected=1,
    )

    result = deidentify_file_copy(source, entity_detector=detector)

    assert result.status == "LOCAL_ONLY"
    audit = json.loads(Path(result.audit_path).read_text(encoding="utf-8"))
    assert audit["entity_assist"]["rejected_candidates"] == 1


def test_docx_entity_assist_handles_name_split_across_runs(tmp_path: Path):
    source = tmp_path / "split-name.docx"
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("张")
    first.bold = True
    second = paragraph.add_run("三在实验室工作")
    second.italic = True
    document.save(source)

    detector = FakeDetector([AssistedEntityLiteral("PERSON", "张三")])
    result = deidentify_file_copy(source, entity_detector=detector)

    rewritten = Document(result.output_path)
    runs = rewritten.paragraphs[0].runs
    text = "".join(run.text for run in runs)
    assert "[PERSON_001]" in text
    assert "张三" not in text
    assert runs[0].bold is True
    assert runs[1].italic is True


def test_persistent_mapper_supports_assisted_entity_types_without_raw_literals():
    secret = b"s" * 32
    mapper = DirectIdentifierMapper(key_secret=secret)
    mapper.register_assisted_literals([AssistedEntityLiteral("PERSON", "张三")])
    assert mapper.replace("张三") == "[PERSON_001]"

    state = mapper.export_state()
    serialized = json.dumps(state, ensure_ascii=False)
    assert "张三" not in serialized

    restored = DirectIdentifierMapper(key_secret=secret, state=state)
    restored.begin_document()
    restored.register_assisted_literals([AssistedEntityLiteral("PERSON", "张三")])
    assert restored.replace("张三") == "[PERSON_001]"


def test_overlapping_semantic_candidate_forces_local_only(tmp_path: Path):
    source = tmp_path / "overlap.txt"
    source.write_text("Contact person@example.com", encoding="utf-8")
    detector = FakeDetector(
        [
            AssistedEntityLiteral(
                "QUASI_IDENTIFIER",
                "Contact person@example.com",
            )
        ]
    )

    result = deidentify_file_copy(source, entity_detector=detector)

    output = Path(result.output_path).read_text(encoding="utf-8")
    assert output == "Contact [EMAIL_001]"
    assert result.status == "LOCAL_ONLY"
    audit = json.loads(Path(result.audit_path).read_text(encoding="utf-8"))
    assert audit["entity_assist"]["overlap_conflicts"] == 1
    assert "person@example.com" not in json.dumps(audit)
