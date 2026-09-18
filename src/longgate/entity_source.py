from __future__ import annotations

import zipfile
from pathlib import Path

_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DOCX_MAX_ENTRIES = 10000
_DOCX_MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024


def _xml_text(data: bytes, *, word_text_only: bool) -> str:
    try:
        from lxml import etree
    except ImportError as exc:
        raise RuntimeError(
            "Structured entity assist for DOCX requires the documents extra."
        ) from exc

    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        huge_tree=False,
        recover=False,
    )
    # Hardened parser: no DTD/entities/network, bounded DOCX package input.
    root = etree.fromstring(data, parser=parser)  # noqa: S320
    if word_text_only:
        return "\n".join(
            node.text or ""
            for node in root.xpath(
                ".//w:t | .//w:delText",
                namespaces={"w": _WORD_NS},
            )
            if node.text
        )
    return "\n".join(value for value in root.itertext() if value)


def _is_word_text_part(name: str) -> bool:
    if name in {
        "word/document.xml",
        "word/comments.xml",
        "word/footnotes.xml",
        "word/endnotes.xml",
    }:
        return True
    if name.startswith("word/header") and name.endswith(".xml"):
        return name[len("word/header") : -4].isdigit()
    if name.startswith("word/footer") and name.endswith(".xml"):
        return name[len("word/footer") : -4].isdigit()
    return False


def _collect_docx(path: Path) -> str:
    try:
        archive = zipfile.ZipFile(path, "r")
    except zipfile.BadZipFile as exc:
        raise ValueError("Input is not a valid DOCX/OOXML ZIP package.") from exc
    parts: list[str] = []
    with archive:
        infos = archive.infolist()
        if len(infos) > _DOCX_MAX_ENTRIES:
            raise ValueError("DOCX contains too many package entries.")
        if sum(info.file_size for info in infos) > _DOCX_MAX_UNCOMPRESSED_BYTES:
            raise ValueError("DOCX uncompressed package size exceeds the safety limit.")
        for info in infos:
            name = info.filename
            if _is_word_text_part(name):
                parts.append(_xml_text(archive.read(name), word_text_only=True))
            elif name in {
                "docProps/core.xml",
                "docProps/app.xml",
                "docProps/custom.xml",
            }:
                parts.append(_xml_text(archive.read(name), word_text_only=False))
    return "\n".join(part for part in parts if part)


def _collect_html(path: Path) -> str:
    try:
        from bs4 import BeautifulSoup, Comment, Doctype
    except ImportError as exc:
        raise RuntimeError(
            "HTML entity assist requires: pip install 'long-gate[documents]'"
        ) from exc
    try:
        source = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError("HTML entity assist currently requires UTF-8 input.") from exc
    soup = BeautifulSoup(source, "html.parser")
    ignored_tags = {"script", "style", "noscript", "template", "svg"}
    parts: list[str] = []
    for node in soup.find_all(string=True):
        if isinstance(node, (Comment, Doctype)):
            continue
        if any(getattr(parent, "name", None) in ignored_tags for parent in node.parents):
            continue
        if str(node).strip():
            parts.append(str(node))
    attribute_names = {
        "alt",
        "title",
        "aria-label",
        "placeholder",
        "value",
        "href",
        "src",
        "action",
    }
    for tag in soup.find_all(True):
        for name in attribute_names:
            value = tag.attrs.get(name)
            if isinstance(value, str) and value.strip():
                parts.append(value)
            elif isinstance(value, list):
                parts.extend(str(item) for item in value if str(item).strip())
    return "\n".join(parts)


def _collect_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False, keep_links=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if (
                    isinstance(value, str)
                    and cell.data_type != "f"
                    and not value.startswith("=")
                    and value.strip()
                ):
                    parts.append(value)
                if cell.comment is not None:
                    if cell.comment.text:
                        parts.append(cell.comment.text)
                    if cell.comment.author:
                        parts.append(cell.comment.author)
                if cell.hyperlink is not None and cell.hyperlink.target:
                    parts.append(str(cell.hyperlink.target))
    for name in (
        "creator",
        "lastModifiedBy",
        "title",
        "subject",
        "description",
        "keywords",
        "category",
    ):
        value = getattr(workbook.properties, name, None)
        if isinstance(value, str) and value.strip():
            parts.append(value)
    return "\n".join(parts)


def collect_entity_assist_text(input_path: str | Path) -> str:
    path = Path(input_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8")
    if suffix in {".html", ".htm"}:
        return _collect_html(path)
    if suffix == ".xlsx":
        return _collect_xlsx(path)
    if suffix == ".docx":
        return _collect_docx(path)
    raise ValueError(
        "Structured entity assist supports TXT/Markdown, HTML, XLSX, and DOCX only."
    )
