from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path

from .pii import scan_text
from .unstructured import (
    SUPPORTED_TEXT,
    load_text_file,
)


class DocumentDependencyMissing(RuntimeError):
    pass


@dataclass(frozen=True)
class DocumentInspection:
    path: str
    kind: str
    characters: int
    text_units: int
    pii_hits: int
    pii_by_entity: dict[str, int]
    release_allowed: bool
    extraction_note: str
    reason: str

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


def _extract_docx(
    path: Path,
) -> tuple[str, int, str]:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentDependencyMissing(
            "DOCX inspection requires the optional dependency: "
            "pip install 'long-gate[documents]'"
        ) from exc

    document = Document(path)
    chunks = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text
    ]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    chunks.append(cell.text)
    return (
        "\n".join(chunks),
        len(chunks),
        (
            "Extracted paragraph and table-cell text locally. "
            "Embedded images, handwriting, and text inside images "
            "are not OCR-scanned."
        ),
    )


def _extract_pdf(
    path: Path,
) -> tuple[str, int, str]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentDependencyMissing(
            "PDF inspection requires the optional dependency: "
            "pip install 'long-gate[documents]'"
        ) from exc

    try:
        reader = PdfReader(path)
        pages = []
        for page in reader.pages:
            pages.append(
                page.extract_text()
                or ""
            )
    except Exception as exc:
        raise RuntimeError(
            "PDF text extraction failed locally. "
            "No network fallback is attempted."
        ) from exc

    return (
        "\n".join(pages),
        len(pages),
        (
            "Text-layer extraction only. Scanned/image-only PDF content "
            "may not be visible because Long Gate does not OCR PDFs "
            "in this baseline."
        ),
    )


def extract_document_text(
    path: str | Path,
) -> tuple[str, str, int, str]:
    document_path = Path(path)
    suffix = document_path.suffix.lower()

    if suffix in SUPPORTED_TEXT:
        text = load_text_file(
            document_path
        )
        return (
            text,
            suffix.lstrip("."),
            text.count("\n")
            + (1 if text else 0),
            "UTF-8 local text extraction.",
        )
    if suffix == ".docx":
        text, units, note = _extract_docx(
            document_path
        )
        return (
            text,
            "docx",
            units,
            note,
        )
    if suffix == ".pdf":
        text, units, note = _extract_pdf(
            document_path
        )
        return (
            text,
            "pdf",
            units,
            note,
        )

    raise ValueError(
        "Unsupported document type: "
        f"{suffix}. Supported: "
        "TXT, Markdown, DOCX, PDF."
    )


def inspect_document_file(
    path: str | Path,
) -> DocumentInspection:
    document_path = Path(path)
    text, kind, units, note = (
        extract_document_text(
            document_path
        )
    )
    findings = scan_text(text)

    return DocumentInspection(
        path=document_path.name,
        kind=kind,
        characters=len(text),
        text_units=units,
        pii_hits=findings.total_hits,
        pii_by_entity=findings.by_entity,
        release_allowed=False,
        extraction_note=note,
        reason=(
            "Document extraction and PII scanning are local-only. "
            "A zero-hit result is not a semantic privacy guarantee "
            "and does not grant network-egress permission."
        ),
    )
