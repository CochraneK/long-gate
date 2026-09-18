from __future__ import annotations

import hashlib
import html
import re
import zipfile
from io import BytesIO
from pathlib import Path

from .entity_source import collect_entity_assist_text
from .format_deidentify import (
    DirectIdentifierMapper,
    FormatPreservingResult,
    SUPPORTED_PRESERVE_TEXT,
    default_output_path,
    deidentify_text_copy,
)
from .pii import scan_text
from .utils import atomic_write_bytes, atomic_write_text, sha256_file, utc_now, write_json

_HTML_SUFFIXES = {".html", ".htm"}
_XLSX_SUFFIXES = {".xlsx"}
_DOCX_SUFFIXES = {".docx"}
_DOCX_MAX_ENTRIES = 10000
_DOCX_MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _paths_and_snapshot(
    input_path: str | Path,
    output_path: str | Path | None,
    allowed_suffixes: set[str],
    *,
    expected_source_sha256: str | None = None,
) -> tuple[Path, Path, bytes, str]:
    source = Path(input_path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    suffix = source.suffix.lower()
    if suffix not in allowed_suffixes:
        raise ValueError(f"Unsupported format-preserving input type: {suffix}")

    destination = (
        Path(output_path).expanduser().resolve()
        if output_path is not None
        else default_output_path(source)
    )
    if destination == source:
        raise ValueError("Refusing to overwrite the input file.")
    if destination.suffix.lower() != suffix:
        raise ValueError("Format-preserving output must keep the same file extension.")

    source_bytes = source.read_bytes()
    digest = hashlib.sha256(source_bytes).hexdigest()
    if expected_source_sha256 is not None and digest != expected_source_sha256:
        raise RuntimeError(
            "Input changed after entity detection; no output was written."
        )
    return source, destination, source_bytes, digest


def _render_report(
    *,
    format_kind: str,
    status: str,
    input_sha256: str,
    output_sha256: str,
    replacements: int,
    entity_counts: dict[str, int],
    remaining_direct_pii_hits: int,
    processed_units: dict[str, int],
    unprocessed_regions: dict[str, int],
    note: str,
    entity_assist: dict[str, object] | None = None,
) -> str:
    entity_rows = "".join(
        f"<tr><td>{html.escape(entity)}</td><td>{count}</td></tr>"
        for entity, count in sorted(entity_counts.items())
    ) or "<tr><td>none</td><td>0</td></tr>"
    processed = "".join(
        f"<li>{html.escape(name)}: {count}</li>"
        for name, count in sorted(processed_units.items())
    ) or "<li>none</li>"
    unprocessed = "".join(
        f"<li>{html.escape(name)}: {count}</li>"
        for name, count in sorted(unprocessed_regions.items())
    ) or "<li>none</li>"
    assist_html = ""
    if entity_assist:
        accepted = int(entity_assist.get("accepted_candidates", 0))
        rejected = int(entity_assist.get("rejected_candidates", 0))
        model_file = html.escape(str(entity_assist.get("model_file", "local model")))
        assist_html = (
            "<div class=\"card\"><h2>本地语义实体辅助</h2>"
            f"<p>模型：<code>{model_file}</code></p>"
            f"<p>已接受候选：{accepted}；被拒绝候选：{rejected}</p>"
            "<p>模型只提名原文 literal；Long Gate 不允许自由改写文件内容。"
            "候选原文不会写入报告。</p></div>"
        )
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>Long Gate 格式保真脱敏报告</title>
<style>
body{{font-family:system-ui,-apple-system,sans-serif;max-width:980px;margin:40px auto;padding:0 20px;line-height:1.6;color:#18212f}}
.card{{border:1px solid #d8dee8;border-radius:14px;padding:18px;margin:16px 0}}
code{{background:#f2f4f8;padding:2px 5px;border-radius:5px;word-break:break-all}}
table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #d8dee8;padding:8px;text-align:left}}th{{background:#f4f7fb}}
</style>
</head>
<body>
<h1>Long Gate · 格式保真脱敏报告</h1>
<div class="card">
<p><strong>状态：{html.escape(status)}</strong></p>
<p>格式：<code>{html.escape(format_kind)}</code></p>
<p>{html.escape(note)}</p>
<p>这不是匿名化证明，也不会自动获得联网外发权限。</p>
</div>
<div class="card">
<h2>完整性</h2>
<p>输入/输出文件名和本地目录路径不会写入本报告。</p>
<p>输入 SHA-256：<code>{html.escape(input_sha256)}</code></p>
<p>输出 SHA-256：<code>{html.escape(output_sha256)}</code></p>
</div>
<div class="card">
<h2>替换</h2>
<p>总替换次数：{replacements}</p>
<table><thead><tr><th>类型</th><th>出现次数</th></tr></thead><tbody>{entity_rows}</tbody></table>
<p>仍检测到的直接 PII：{remaining_direct_pii_hits}</p>
</div>
<div class="card"><h2>已处理区域</h2><ul>{processed}</ul></div>
<div class="card"><h2>未自动改写 / 需复核区域</h2><ul>{unprocessed}</ul></div>
{assist_html}
<div class="card">
<h2>人工复核重点</h2>
<ul>
<li>姓名、组织、地点、别名、罕见事件和组合身份线索仍需复核。</li>
<li>结构保真不等于语义匿名；任何 remaining direct PII 都会保持 LOCAL_ONLY。</li>
<li>脱敏质量判断与网络外发授权相互独立。</li>
</ul>
</div>
<p>报告由 Long Gate 于 {html.escape(utc_now())} 在本地生成，不嵌入原文、脱敏正文、文件名或本地路径。</p>
</body></html>"""


def _commit_result(
    *,
    source: Path,
    destination: Path,
    source_sha256: str,
    payload: bytes,
    format_kind: str,
    mapper: DirectIdentifierMapper,
    remaining_direct_pii_hits: int,
    remaining_by_entity: dict[str, int],
    processed_units: dict[str, int],
    unprocessed_regions: dict[str, int],
    note: str,
    force_local_only: bool = False,
    entity_assist: dict[str, object] | None = None,
) -> FormatPreservingResult:
    if sha256_file(source) != source_sha256:
        raise RuntimeError("Input file changed during processing; no output was written.")

    atomic_write_bytes(destination, payload)
    if sha256_file(source) != source_sha256:
        destination.unlink(missing_ok=True)
        raise RuntimeError(
            "Input file changed while output was being committed; generated output was removed."
        )

    output_sha256 = sha256_file(destination)
    assist_rejected = int((entity_assist or {}).get("rejected_candidates", 0))
    status = (
        "LOCAL_ONLY"
        if remaining_direct_pii_hits or force_local_only or assist_rejected
        else "MANUAL_REVIEW_REQUIRED"
    )
    audit_path = destination.with_name(destination.name + ".audit.json")
    report_path = destination.with_name(destination.name + ".trust-report.html")
    next_actions = [
        "人工复核姓名、组织、地点、别名、罕见事件与组合身份线索。",
        "检查报告中的未自动改写区域；任何剩余直接 PII 都必须保持 LOCAL_ONLY。",
        "不要把格式保真处理结果视为自动联网授权。",
    ]
    payload_json = {
        "format": "long-gate-format-preserving-deidentify-v2",
        "status": status,
        "input_format": format_kind,
        "input_sha256": source_sha256,
        "output_sha256": output_sha256,
        "replacements": mapper.replacements,
        "replacement_entities": dict(mapper.entity_counts),
        "remaining_direct_pii_hits": remaining_direct_pii_hits,
        "remaining_direct_pii_by_entity": remaining_by_entity,
        "processed_units": processed_units,
        "unprocessed_regions": unprocessed_regions,
        "original_unchanged": True,
        "format_preserved": True,
        "manual_review_required": True,
        "automatic_release_allowed": False,
        "release_allowed": False,
        "next_actions": next_actions,
        "entity_assist": entity_assist,
    }
    write_json(audit_path, payload_json)
    atomic_write_text(
        report_path,
        _render_report(
            format_kind=format_kind,
            status=status,
            input_sha256=source_sha256,
            output_sha256=output_sha256,
            replacements=mapper.replacements,
            entity_counts=dict(mapper.entity_counts),
            remaining_direct_pii_hits=remaining_direct_pii_hits,
            processed_units=processed_units,
            unprocessed_regions=unprocessed_regions,
            note=note,
            entity_assist=entity_assist,
        ),
    )
    return FormatPreservingResult(
        status=status,
        input_path=str(source),
        output_path=str(destination),
        audit_path=str(audit_path),
        report_path=str(report_path),
        input_sha256=source_sha256,
        output_sha256=output_sha256,
        replacements=mapper.replacements,
        replacement_entities=dict(mapper.entity_counts),
        original_unchanged=True,
        manual_review_required=True,
        automatic_release_allowed=False,
        release_allowed=False,
        next_actions=next_actions,
        entity_assist=entity_assist,
    )


def deidentify_html_copy(
    input_path: str | Path,
    output_path: str | Path | None = None,
    *,
    mapper: DirectIdentifierMapper | None = None,
    entity_assist: dict[str, object] | None = None,
    expected_source_sha256: str | None = None,
) -> FormatPreservingResult:
    try:
        from bs4 import BeautifulSoup, Comment, Doctype, NavigableString
    except ImportError as exc:
        raise RuntimeError(
            "HTML format-preserving de-identification requires: "
            "pip install 'long-gate[documents]'"
        ) from exc

    source, destination, source_bytes, source_sha256 = _paths_and_snapshot(
        input_path,
        output_path,
        _HTML_SUFFIXES,
        expected_source_sha256=expected_source_sha256,
    )
    try:
        source_text = source_bytes.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError("HTML format-preserving path currently requires UTF-8 input.") from exc

    soup = BeautifulSoup(source_text, "html.parser")
    mapper = mapper or DirectIdentifierMapper()
    ignored_tags = {"script", "style", "noscript", "template", "svg"}
    processed_text_nodes = 0
    ignored_text_nodes = 0

    for node in list(soup.find_all(string=True)):
        if isinstance(node, (Comment, Doctype)):
            ignored_text_nodes += 1
            continue
        if any(getattr(parent, "name", None) in ignored_tags for parent in node.parents):
            ignored_text_nodes += 1
            continue
        transformed = mapper.replace(str(node))
        if transformed != str(node):
            node.replace_with(NavigableString(transformed))
        processed_text_nodes += 1

    processed_attributes = 0
    attribute_names = {
        "alt",
        "title",
        "aria-label",
        "placeholder",
        "value",
        "href",
        "src",
        "action",
    }
    for tag in soup.find_all(True):
        for name in list(tag.attrs):
            if name not in attribute_names:
                continue
            value = tag.attrs.get(name)
            if isinstance(value, str):
                tag.attrs[name] = mapper.replace(value)
                processed_attributes += 1
            elif isinstance(value, list):
                tag.attrs[name] = [mapper.replace(str(item)) for item in value]
                processed_attributes += len(value)

    output_text = str(soup)
    remaining = scan_text(output_text)
    return _commit_result(
        source=source,
        destination=destination,
        source_sha256=source_sha256,
        payload=output_text.encode("utf-8"),
        format_kind="html",
        mapper=mapper,
        remaining_direct_pii_hits=remaining.total_hits,
        remaining_by_entity=remaining.by_entity,
        processed_units={
            "visible_text_nodes": processed_text_nodes,
            "selected_attributes": processed_attributes,
        },
        unprocessed_regions={
            "comments_doctype_or_ignored_text_nodes": ignored_text_nodes,
        },
        note=(
            "HTML DOM and tag structure are retained, but BeautifulSoup may normalize "
            "source serialization/whitespace. Visible text and selected user-facing/link "
            "attributes are processed; scripts/styles/templates/SVG text are not rewritten."
        ),
        entity_assist=entity_assist,
    )


def _add_findings(
    text: str,
    total: int,
    by_entity: dict[str, int],
) -> tuple[int, dict[str, int]]:
    findings = scan_text(text)
    total += findings.total_hits
    for entity, count in findings.by_entity.items():
        by_entity[entity] = by_entity.get(entity, 0) + count
    return total, by_entity


def deidentify_xlsx_copy(
    input_path: str | Path,
    output_path: str | Path | None = None,
    *,
    mapper: DirectIdentifierMapper | None = None,
    entity_assist: dict[str, object] | None = None,
    expected_source_sha256: str | None = None,
) -> FormatPreservingResult:
    from openpyxl import load_workbook

    source, destination, source_bytes, source_sha256 = _paths_and_snapshot(
        input_path,
        output_path,
        _XLSX_SUFFIXES,
        expected_source_sha256=expected_source_sha256,
    )
    workbook = load_workbook(BytesIO(source_bytes), data_only=False, keep_links=True)
    mapper = mapper or DirectIdentifierMapper()
    processed_cells = 0
    processed_comments = 0
    processed_hyperlinks = 0
    processed_properties = 0
    formulas_unmodified = 0
    sheet_titles_unmodified = 0
    defined_names_unmodified = 0
    remaining_total = 0
    remaining_by_entity: dict[str, int] = {}

    for sheet in workbook.worksheets:
        title_findings = scan_text(sheet.title)
        if title_findings.total_hits:
            sheet_titles_unmodified += 1
            remaining_total += title_findings.total_hits
            for entity, count in title_findings.by_entity.items():
                remaining_by_entity[entity] = remaining_by_entity.get(entity, 0) + count

        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str):
                    if cell.data_type == "f" or value.startswith("="):
                        formulas_unmodified += 1
                        remaining_total, remaining_by_entity = _add_findings(
                            value, remaining_total, remaining_by_entity
                        )
                    else:
                        cell.value = mapper.replace(value)
                        processed_cells += 1

                if cell.comment is not None:
                    cell.comment.text = mapper.replace(cell.comment.text or "")
                    cell.comment.author = mapper.replace(cell.comment.author or "")
                    processed_comments += 1

                if cell.hyperlink is not None and cell.hyperlink.target:
                    cell.hyperlink.target = mapper.replace(str(cell.hyperlink.target))
                    processed_hyperlinks += 1

    for name in (
        "creator",
        "lastModifiedBy",
        "title",
        "subject",
        "description",
        "keywords",
        "category",
    ):
        value = getattr(workbook.properties, name, None)
        if isinstance(value, str) and value:
            transformed = mapper.replace(value)
            setattr(workbook.properties, name, transformed)
            processed_properties += 1
            remaining_total, remaining_by_entity = _add_findings(
                transformed, remaining_total, remaining_by_entity
            )

    for defined_name in workbook.defined_names.values():
        value = getattr(defined_name, "attr_text", None)
        if isinstance(value, str) and value:
            findings = scan_text(value)
            if findings.total_hits:
                defined_names_unmodified += 1
                remaining_total += findings.total_hits
                for entity, count in findings.by_entity.items():
                    remaining_by_entity[entity] = remaining_by_entity.get(entity, 0) + count

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                for value in (
                    cell.value,
                    cell.comment.text if cell.comment is not None else None,
                    cell.comment.author if cell.comment is not None else None,
                    cell.hyperlink.target if cell.hyperlink is not None else None,
                ):
                    if isinstance(value, str) and not (
                        cell.data_type == "f" and value == cell.value
                    ):
                        remaining_total, remaining_by_entity = _add_findings(
                            value, remaining_total, remaining_by_entity
                        )

    out = BytesIO()
    workbook.save(out)
    return _commit_result(
        source=source,
        destination=destination,
        source_sha256=source_sha256,
        payload=out.getvalue(),
        format_kind="xlsx",
        mapper=mapper,
        remaining_direct_pii_hits=remaining_total,
        remaining_by_entity=remaining_by_entity,
        processed_units={
            "worksheets": len(workbook.worksheets),
            "string_cells": processed_cells,
            "comments": processed_comments,
            "hyperlink_targets": processed_hyperlinks,
            "workbook_properties": processed_properties,
        },
        unprocessed_regions={
            "formula_cells": formulas_unmodified,
            "sheet_titles_with_direct_pii": sheet_titles_unmodified,
            "defined_names_with_direct_pii": defined_names_unmodified,
        },
        note=(
            "All worksheets, including hidden sheets, are traversed. String cells, comments, "
            "hyperlink targets, and selected workbook properties are processed. Formulas and "
            "sheet titles/defined names are preserved and reported for manual review."
        ),
        entity_assist=entity_assist,
    )



def _set_ooxml_text(node: object, value: str) -> None:
    node.text = value
    if value.startswith(" ") or value.endswith(" "):
        node.set(_XML_SPACE, "preserve")


def _node_offsets(nodes: list[object]) -> list[tuple[int, int]]:
    offsets: list[tuple[int, int]] = []
    cursor = 0
    for node in nodes:
        text = node.text or ""
        offsets.append((cursor, cursor + len(text)))
        cursor += len(text)
    return offsets


def _locate_start(offsets: list[tuple[int, int]], position: int) -> tuple[int, int]:
    for index, (start, end) in enumerate(offsets):
        if start <= position < end:
            return index, position - start
    raise ValueError("Identifier start did not map to an OOXML text node.")


def _locate_end(offsets: list[tuple[int, int]], position: int) -> tuple[int, int]:
    for index, (start, end) in enumerate(offsets):
        if start < position <= end:
            return index, position - start
    raise ValueError("Identifier end did not map to an OOXML text node.")


def _rewrite_ooxml_text_nodes(nodes: list[object], mapper: DirectIdentifierMapper) -> int:
    if not nodes:
        return 0
    original_texts = [node.text or "" for node in nodes]
    joined = "".join(original_texts)
    if not joined:
        return 0
    plan = mapper.plan(joined)
    if not plan:
        return 0

    offsets = _node_offsets(nodes)
    for span in reversed(plan):
        start_index, start_local = _locate_start(offsets, span.start)
        end_index, end_local = _locate_end(offsets, span.end)

        if start_index == end_index:
            current = nodes[start_index].text or ""
            _set_ooxml_text(
                nodes[start_index],
                current[:start_local] + span.label + current[end_local:],
            )
            continue

        start_text = nodes[start_index].text or ""
        end_text = nodes[end_index].text or ""
        _set_ooxml_text(nodes[start_index], start_text[:start_local] + span.label)
        for index in range(start_index + 1, end_index):
            _set_ooxml_text(nodes[index], "")
        _set_ooxml_text(nodes[end_index], end_text[end_local:])
    return len(plan)


def _scan_word_xml_root(root: object) -> tuple[int, dict[str, int], int]:
    total = 0
    by_entity: dict[str, int] = {}
    paragraphs = root.xpath(".//w:p", namespaces={"w": _WORD_NS})
    for paragraph in paragraphs:
        text = "".join(
            node.text or ""
            for node in paragraph.xpath(".//w:t | .//w:delText", namespaces={"w": _WORD_NS})
        )
        total, by_entity = _add_findings(text, total, by_entity)

    instruction_hits = 0
    for node in root.xpath(".//w:instrText", namespaces={"w": _WORD_NS}):
        if node.text:
            findings = scan_text(node.text)
            instruction_hits += findings.total_hits
            total += findings.total_hits
            for entity, count in findings.by_entity.items():
                by_entity[entity] = by_entity.get(entity, 0) + count
    return total, by_entity, instruction_hits


def _rewrite_word_xml(
    data: bytes,
    mapper: DirectIdentifierMapper,
) -> tuple[bytes, int, int, dict[str, int], int]:
    try:
        from lxml import etree
    except ImportError as exc:
        raise RuntimeError(
            "DOCX format-preserving de-identification requires the documents extra."
        ) from exc

    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        huge_tree=False,
        recover=False,
    )
    # Hardened parser: no DTD/entities/network, bounded OOXML ZIP input.
    root = etree.fromstring(data, parser=parser)  # noqa: S320
    paragraphs = root.xpath(".//w:p", namespaces={"w": _WORD_NS})
    processed_paragraphs = 0
    for paragraph in paragraphs:
        nodes = list(
            paragraph.xpath(".//w:t | .//w:delText", namespaces={"w": _WORD_NS})
        )
        if not nodes:
            continue
        _rewrite_ooxml_text_nodes(nodes, mapper)
        processed_paragraphs += 1

    remaining_total, remaining_by_entity, instruction_hits = _scan_word_xml_root(root)
    rendered = etree.tostring(
        root,
        xml_declaration=data.lstrip().startswith(b"<?xml"),
        encoding="UTF-8",
        standalone=None,
    )
    return (
        rendered,
        processed_paragraphs,
        remaining_total,
        remaining_by_entity,
        instruction_hits,
    )


def _rewrite_property_xml(
    data: bytes,
    mapper: DirectIdentifierMapper,
) -> tuple[bytes, int, int, dict[str, int]]:
    try:
        from lxml import etree
    except ImportError as exc:
        raise RuntimeError(
            "DOCX format-preserving de-identification requires the documents extra."
        ) from exc

    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        huge_tree=False,
        recover=False,
    )
    # Hardened parser: no DTD/entities/network, bounded OOXML ZIP input.
    root = etree.fromstring(data, parser=parser)  # noqa: S320
    processed = 0
    for element in root.iter():
        if element.text:
            transformed = mapper.replace(element.text)
            if transformed != element.text:
                element.text = transformed
                processed += 1
    text = "\n".join(value for value in root.itertext() if value)
    remaining = scan_text(text)
    rendered = etree.tostring(
        root,
        xml_declaration=data.lstrip().startswith(b"<?xml"),
        encoding="UTF-8",
        standalone=None,
    )
    return rendered, processed, remaining.total_hits, remaining.by_entity


def _scan_unhandled_xml(data: bytes) -> tuple[int, dict[str, int]]:
    """Scan XML text plus obviously user-authored string attributes, not structural numbers."""
    try:
        from lxml import etree
    except ImportError as exc:
        raise RuntimeError(
            "DOCX format-preserving de-identification requires the documents extra."
        ) from exc

    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        huge_tree=False,
        recover=False,
    )
    # Hardened parser: no DTD/entities/network, bounded OOXML ZIP input.
    root = etree.fromstring(data, parser=parser)  # noqa: S320
    total = 0
    by_entity: dict[str, int] = {}

    for value in root.itertext():
        if value:
            total, by_entity = _add_findings(value, total, by_entity)

    for element in root.iter():
        for value in element.attrib.values():
            if not isinstance(value, str):
                continue
            lowered = value.lower()
            if "@" in value or "mailto:" in lowered or "tel:" in lowered:
                total, by_entity = _add_findings(value, total, by_entity)

    return total, by_entity


def _is_docx_word_text_part(name: str) -> bool:
    if name in {
        "word/document.xml",
        "word/comments.xml",
        "word/footnotes.xml",
        "word/endnotes.xml",
    }:
        return True
    return bool(re.fullmatch(r"word/(?:header|footer)\d+\.xml", name))


def deidentify_docx_copy(
    input_path: str | Path,
    output_path: str | Path | None = None,
    *,
    mapper: DirectIdentifierMapper | None = None,
    entity_assist: dict[str, object] | None = None,
    expected_source_sha256: str | None = None,
) -> FormatPreservingResult:
    source, destination, source_bytes, source_sha256 = _paths_and_snapshot(
        input_path,
        output_path,
        _DOCX_SUFFIXES,
        expected_source_sha256=expected_source_sha256,
    )
    mapper = mapper or DirectIdentifierMapper()
    remaining_total = 0
    remaining_by_entity: dict[str, int] = {}
    processed_parts = 0
    processed_paragraphs = 0
    processed_property_fields = 0
    relationship_parts_with_pii = 0
    field_instruction_pii_hits = 0
    unprocessed_xml_parts_with_pii = 0
    media_entries = 0
    embedded_entries = 0
    other_risky_binary_entries = 0
    force_unparsed_xml_local_only = False

    input_buffer = BytesIO(source_bytes)
    output_buffer = BytesIO()
    try:
        archive = zipfile.ZipFile(input_buffer, "r")
    except zipfile.BadZipFile as exc:
        raise ValueError("Input is not a valid DOCX/OOXML ZIP package.") from exc

    with archive:
        infos = archive.infolist()
        if any(info.filename.startswith("_xmlsignatures/") for info in infos):
            raise ValueError(
                "Digitally signed DOCX packages are not rewritten because modification "
                "would invalidate the package signature."
            )
        if len(infos) > _DOCX_MAX_ENTRIES:
            raise ValueError("DOCX contains too many package entries.")
        total_uncompressed = sum(info.file_size for info in infos)
        if total_uncompressed > _DOCX_MAX_UNCOMPRESSED_BYTES:
            raise ValueError("DOCX uncompressed package size exceeds the safety limit.")

        with zipfile.ZipFile(output_buffer, "w") as output_archive:
            output_archive.comment = archive.comment
            for info in infos:
                data = archive.read(info.filename)
                name = info.filename
                rewritten = data
                handled_xml = False

                if _is_docx_word_text_part(name):
                    handled_xml = True
                    (
                        rewritten,
                        paragraph_count,
                        part_remaining,
                        part_by_entity,
                        instruction_hits,
                    ) = _rewrite_word_xml(data, mapper)
                    processed_parts += 1
                    processed_paragraphs += paragraph_count
                    field_instruction_pii_hits += instruction_hits
                    remaining_total += part_remaining
                    for entity, count in part_by_entity.items():
                        remaining_by_entity[entity] = (
                            remaining_by_entity.get(entity, 0) + count
                        )
                elif name in {
                    "docProps/core.xml",
                    "docProps/app.xml",
                    "docProps/custom.xml",
                }:
                    handled_xml = True
                    (
                        rewritten,
                        field_count,
                        part_remaining,
                        part_by_entity,
                    ) = _rewrite_property_xml(data, mapper)
                    processed_property_fields += field_count
                    remaining_total += part_remaining
                    for entity, count in part_by_entity.items():
                        remaining_by_entity[entity] = (
                            remaining_by_entity.get(entity, 0) + count
                        )
                elif name.endswith(".rels"):
                    handled_xml = True
                    try:
                        relation_text = data.decode("utf-8")
                    except UnicodeDecodeError:
                        relation_text = ""
                    findings = scan_text(relation_text)
                    if findings.total_hits:
                        relationship_parts_with_pii += 1
                        remaining_total += findings.total_hits
                        for entity, count in findings.by_entity.items():
                            remaining_by_entity[entity] = (
                                remaining_by_entity.get(entity, 0) + count
                            )

                if name.endswith(".xml") and not handled_xml:
                    try:
                        part_remaining, part_by_entity = _scan_unhandled_xml(data)
                    except Exception:
                        unprocessed_xml_parts_with_pii += 1
                        force_unparsed_xml_local_only = True
                    else:
                        if part_remaining:
                            unprocessed_xml_parts_with_pii += 1
                            remaining_total += part_remaining
                            for entity, count in part_by_entity.items():
                                remaining_by_entity[entity] = (
                                    remaining_by_entity.get(entity, 0) + count
                                )

                if name.startswith("word/media/") and not name.endswith("/"):
                    media_entries += 1
                if name.startswith("word/embeddings/") and not name.endswith("/"):
                    embedded_entries += 1
                if name.startswith("word/activeX/") and not name.endswith("/"):
                    other_risky_binary_entries += 1
                if (
                    name.startswith("customXml/")
                    and not name.endswith("/")
                    and not name.endswith(".xml")
                    and not name.endswith(".rels")
                ):
                    other_risky_binary_entries += 1

                output_archive.writestr(info, rewritten)

    force_local_only = any(
        (
            media_entries,
            embedded_entries,
            other_risky_binary_entries,
            relationship_parts_with_pii,
            field_instruction_pii_hits,
            unprocessed_xml_parts_with_pii,
            force_unparsed_xml_local_only,
        )
    )
    return _commit_result(
        source=source,
        destination=destination,
        source_sha256=source_sha256,
        payload=output_buffer.getvalue(),
        format_kind="docx",
        mapper=mapper,
        remaining_direct_pii_hits=remaining_total,
        remaining_by_entity=remaining_by_entity,
        processed_units={
            "word_xml_parts": processed_parts,
            "paragraphs": processed_paragraphs,
            "property_fields_replaced": processed_property_fields,
        },
        unprocessed_regions={
            "relationship_parts_with_direct_pii": relationship_parts_with_pii,
            "field_instruction_direct_pii_hits": field_instruction_pii_hits,
            "unprocessed_xml_parts_with_direct_pii": unprocessed_xml_parts_with_pii,
            "unparsed_xml_parts": int(force_unparsed_xml_local_only),
            "media_entries": media_entries,
            "embedded_entries": embedded_entries,
            "other_risky_binary_entries": other_risky_binary_entries,
        },
        note=(
            "DOCX is rewritten at OOXML text-node level. Paragraph/run/table/header/footer/"
            "comment/footnote/endnote package structure is retained, including identifiers "
            "split across runs. Hyperlink relationship targets are not rewritten; images, "
            "embedded objects, ActiveX, and non-XML custom parts are treated as unresolved "
            "local-only risk; parseable custom XML is residual-scanned instead."
        ),
        force_local_only=force_local_only,
        entity_assist=entity_assist,
    )

def deidentify_file_copy(
    input_path: str | Path,
    output_path: str | Path | None = None,
    *,
    mapper: DirectIdentifierMapper | None = None,
    entity_detector: object | None = None,
    entity_max_tokens: int = 768,
) -> FormatPreservingResult:
    shared_mapper = mapper is not None
    mapper = mapper or DirectIdentifierMapper()
    if shared_mapper:
        mapper.begin_document()

    entity_assist: dict[str, object] | None = None
    expected_source_sha256: str | None = None
    if entity_detector is not None:
        expected_source_sha256 = sha256_file(input_path)
        source_for_detection = collect_entity_assist_text(input_path)
        if sha256_file(input_path) != expected_source_sha256:
            raise RuntimeError(
                "Input changed during entity detection extraction; no output was written."
            )
        detect = getattr(entity_detector, "detect", None)
        if not callable(detect):
            raise TypeError("entity_detector must provide a callable detect method.")
        assist_result = detect(source_for_detection, max_tokens=entity_max_tokens)
        candidates = getattr(assist_result, "candidates", None)
        summary = getattr(assist_result, "summary_dict", None)
        if not isinstance(candidates, list) or not callable(summary):
            raise RuntimeError("Entity detector returned an invalid structured result.")
        mapper.register_assisted_literals(candidates)
        entity_assist = summary()

    suffix = Path(input_path).suffix.lower()
    if suffix in SUPPORTED_PRESERVE_TEXT:
        return deidentify_text_copy(
            input_path,
            output_path,
            mapper=mapper,
            entity_assist=entity_assist,
            expected_source_sha256=expected_source_sha256,
        )
    if suffix in _HTML_SUFFIXES:
        return deidentify_html_copy(
            input_path,
            output_path,
            mapper=mapper,
            entity_assist=entity_assist,
            expected_source_sha256=expected_source_sha256,
        )
    if suffix in _XLSX_SUFFIXES:
        return deidentify_xlsx_copy(
            input_path,
            output_path,
            mapper=mapper,
            entity_assist=entity_assist,
            expected_source_sha256=expected_source_sha256,
        )
    if suffix in _DOCX_SUFFIXES:
        return deidentify_docx_copy(
            input_path,
            output_path,
            mapper=mapper,
            entity_assist=entity_assist,
            expected_source_sha256=expected_source_sha256,
        )
    raise ValueError(
        "Format-preserving deidentify supports TXT/Markdown, HTML, XLSX, and DOCX in the "
        "current phase. PDF is not silently flattened; use semantic-summarize "
        "only when an abstract text output is actually intended."
    )
